# AgentCode/create_template.py

import openpyxl, os
base = os.path.dirname(__file__)
p = os.path.join(base, "Templates", "OutputFormat.xlsx")
os.makedirs(os.path.dirname(p), exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Template"
labels = [
    "Market Share",
    "Products Manufactured",
    "Key Raw Materials",
    "Market Cap",
    "Shareholding Ratio",
    "Key Financial Ratio",
    "Key Balance sheet values",
    "Key P&L values",
    "Key Forensic analysis notings",
    "Niche / Differentiator",
    "Legal action pending / issues"
]
ws["A6"] = "Field"
for i, lab in enumerate(labels, start=7):
    ws.cell(row=i, column=1).value = lab
wb.save(p)
print("Template created at", p)